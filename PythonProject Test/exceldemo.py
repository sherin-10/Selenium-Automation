import openpyxl
book = openpyxl.load_workbook("D:\\Automation & Testing\\pythondemo.xlsx") # to load the excel sheet
sheet = book.active # used to fetch the active sheet in the excel
cell = sheet.cell(row= 1, column = 2)
print(cell.value)
#Assigning value back to cell in sheet
sheet.cell(row= 2, column = 2).value =" Rahul"
print(sheet.cell(row= 2, column = 2).value)
print(sheet.max_row)
print(sheet.max_column)
# Another way to access cell value
print(sheet['A5'].value) # prints A5 value in sheet
Dict={}
# Nested for loop to access all elements of the sheet
for i in range(1,sheet.max_row+1): # to get rows
    # if sheet.cell(row=i,column=1).value =="tc3": # if condition to select only testcase3
      for j in range(1, sheet.max_column+1): # to get columns
        #print(sheet.cell(row=i,column=j).value) #print values
        #stores the value to a dictionary
        #each iteration the dictionary is getting updated with new values
        Dict[sheet.cell(row=1,column=j).value]= sheet.cell(row=i,column=j).value
print(Dict)  # only last record is stored and printed
